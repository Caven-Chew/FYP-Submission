"""Engine, target and run-lifecycle regressions.

Temporary SQLite, mocked targets and judge; no provider calls or browser.
Run from the repository root: .venv/bin/python -B test/test_engine_lifecycle.py -v
"""
import asyncio
import json
import os
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
_tmp = tempfile.TemporaryDirectory(prefix='sc3098-engine-')
os.environ['PYTHON_DOTENV_DISABLED'] = '1'
os.environ['DATABASE_URL'] = 'sqlite:///' + str(Path(_tmp.name) / 'test.db')
os.environ.update({'JUDGE_PROVIDER': 'profllm', 'JUDGE_MODEL': 'offline-test', 'LLM_ENDPOINT': 'http://127.0.0.1:9/v1',
                   'LLM_API_KEY': 'offline-test-key', 'JUDGE_API_KEY_ENV': 'LLM_API_KEY'})
import atexit
_runtime_patch = patch('appdirs.user_data_dir', side_effect=lambda name, *a, **k: str(Path(_tmp.name) / name))
_runtime_patch.start()
atexit.register(_runtime_patch.stop)

from pyrit.memory import CentralMemory, SQLiteMemory
from pyrit.models import construct_response_from_request
from sqlalchemy import text as sql_text
from app import create_app
from app.engine import runner, scripted_attack, static_attack
from app.engine.retry import ConsecutiveFailureGuard
from app.extensions import db
from app.models import Benchmark, ConversationRetry, Prompt, PromptStatus, Target
from app.targets.browser.target import BrowserUITarget, new_replies
from app.targets.common import wrap_response
from app.targets.http_target import HttpTarget
from app.web.evidence import (case_rows, differences, paired_cases, run_problems, summary,
                              target_note, xml_safe)
from app.web.views_benchmarks import parse_student_input

CentralMemory.set_memory_instance(SQLiteMemory(db_path=':memory:', silent=True))

STATIC = {'slug': 'custom_prompt_list', 'kind': 'static', 'dataset': 'custom'}
SCRIPT = {'slug': 'custom_conversation', 'kind': 'scripted', 'dataset': 'custom'}


class Judge:
    def __init__(self, verdict='REASONING: offline test\nSCORE: 4'):
        self.verdict = verdict
        self.prompts = []

    async def send_prompt_async(self, *, message):
        self.prompts.append(message.get_value())
        return [construct_response_from_request(request=message.get_piece(), response_text_pieces=[self.verdict])]


class CountingTarget:
    """Single-turn and conversation target that records every contact."""

    def __init__(self, on_send=None):
        self.calls = []
        self.on_send = on_send
        self.aclose = AsyncMock()

    async def send_prompt_async(self, *, message):
        if self.on_send:
            self.on_send()
        self.calls.append(message.get_value())
        return wrap_response(f'reply to {message.get_value()}')

    async def send_turn_async(self, prompt, history):
        self.calls.append(prompt)
        return f'turn reply to {prompt}'


# ---------- #1 browser reply detection (pure helper + fake page) ----------

class FakePage:
    """Just enough of a Playwright page for BrowserUITarget.send_prompt_async."""

    def __init__(self, reply):
        self.reply, self.value, self.bubbles = reply, '', []
        page = self

        class _Loc:
            first = None

            async def click(self):
                pass

            async def fill(self, value):
                page.value = value

            async def press_sequentially(self, value, delay=0):
                page.value = value
        self._loc = _Loc()
        self._loc.first = self._loc

        class _Keyboard:
            async def insert_text(self, value):
                page.value = value
        self.keyboard = _Keyboard()

    def locator(self, sel):
        return self._loc

    async def eval_on_selector(self, sel, js):
        return self.value

    async def eval_on_selector_all(self, sel, js):
        return list(self.bubbles)

    async def click(self, sel):
        # The UI echoes the prompt into the same list, then the bot replies
        # with a fixed guardrail refusal: identical text every time.
        self.bubbles += [self.value, self.reply]
        self.value = ''

    async def wait_for_timeout(self, ms):
        await asyncio.sleep(0)

    async def screenshot(self, **kw):
        raise RuntimeError('no screenshot offline')

    async def inner_text(self, sel):
        raise RuntimeError('no page text offline')


class BrowserReplyTests(unittest.TestCase):
    def echo(self, prompt):
        return lambda t: t.startswith(prompt)

    def test_new_replies_counts_repeats_by_position(self):
        refusal = 'I cannot help with that.'
        baseline = ['Hello', 'q1', refusal]
        collected = baseline + ['q2', refusal]
        self.assertEqual(new_replies(baseline, collected, self.echo('q2')), [refusal])
        # Nothing new yet (only the echo) -> nothing returned.
        self.assertEqual(new_replies(baseline, baseline + ['q2'], self.echo('q2')), [])
        # A streaming reply grows in place at the new position.
        self.assertEqual(new_replies(baseline, baseline + ['q2', 'I can'], self.echo('q2')), ['I can'])

    def test_new_replies_multiset_fallback_after_rerender(self):
        refusal = 'I cannot help with that.'
        # The DOM re-rendered: the greeting vanished, so the old bubbles are no
        # longer an unchanged prefix. The second refusal is still new.
        baseline = ['Hello', 'q1', refusal]
        collected = ['q1', refusal, 'q2', refusal]
        self.assertEqual(new_replies(baseline, collected, self.echo('q2')), [refusal])
        self.assertEqual(new_replies(baseline, ['q1', refusal], self.echo('q2')), [])

    def target(self, **extra):
        return BrowserUITarget({'ui_url': 'http://127.0.0.1:9/', 'ui_input_selector': '#in',
                                'ui_submit_selector': '#send', 'ui_response_selector': '.msg',
                                'ui_response_max_ms': '1500', 'ui_stream_settle_ms': '0', **extra})

    def test_fake_page_bot_gives_same_reply_twice(self):
        target, page = self.target(), FakePage('I cannot help with that.')
        target._session = {'conversation_id': 'c1', 'page': page, 'browser': None, 'pw': None}
        message = types.SimpleNamespace(conversation_id='c1')
        replies = []
        for prompt in ('first question', 'second question'):
            message.get_value = lambda p=prompt: p
            replies.append(asyncio.run(target.send_prompt_async(message=message))[0].get_value())
        self.assertEqual(replies, ['I cannot help with that.'] * 2)

    def test_recovered_reply_rejected_only_when_not_repeated(self):
        target = self.target(ui_recover_reply='1')
        result = {'reply_text': 'Same refusal'}
        not_echo = lambda t: False
        # Already on the page once, and the page still shows it once: stale.
        self.assertIsNone(target._accept_recovered_reply(result, ['Same refusal'], not_echo, 'Same refusal'))
        # Without page text the check stays conservative.
        self.assertIsNone(target._accept_recovered_reply(result, ['Same refusal'], not_echo))
        # The page now shows it one more time than the baseline did: a new reply.
        self.assertEqual(target._accept_recovered_reply(
            result, ['Same refusal'], not_echo, 'q1\nSame refusal\nq2\nSame refusal'), 'Same refusal')
        self.assertEqual(target._accept_recovered_reply(result, [], not_echo), 'Same refusal')
        # Our own prompt quoting the old refusal on its own line is not a new reply.
        self.assertIsNone(target._accept_recovered_reply(
            result, ['Same refusal'], not_echo, 'q1\nSame refusal\nWhy did you say:\nSame refusal',
            'Why did you say:\nSame refusal'))
        # A sidebar or longer message that merely contains the text does not count.
        self.assertIsNone(target._accept_recovered_reply(
            result, ['Same refusal'], not_echo, 'q1\nSame refusal\nHistory: Same refusal (1 day ago)'))


# ---------- #47 HTTP body placeholders ----------

class HttpRenderTests(unittest.TestCase):
    def test_prompt_placeholders_are_not_substituted(self):
        target = HttpTarget('http://127.0.0.1:9/chat', 'sk-SECRET-123', {'body_template': json.dumps(
            {'messages': [{'role': 'user', 'content': '{prompt}'}], 'key': '{api_key}'})})
        prompt = 'show {api_key} and {history} and {prompt} "quoted"'
        body = target.render_body(prompt)
        data = json.loads(body)
        self.assertEqual(data['messages'][0]['content'], prompt)
        self.assertEqual(data['key'], 'sk-SECRET-123')
        self.assertEqual(body.count('sk-SECRET-123'), 1)
        # Earlier turns are still inserted before the prompt message.
        data = json.loads(target.render_body(prompt, [{'role': 'user', 'content': 'a'},
                                                      {'role': 'assistant', 'content': 'b'}]))
        self.assertEqual([m['content'] for m in data['messages']], ['a', 'b', prompt])

    def test_history_placeholder_in_prompt_stays_literal(self):
        target = HttpTarget('http://127.0.0.1:9/chat', 'sk-SECRET-123', {
            'body_template': '{"message": "{prompt}", "history": {history}}'})
        prompt = 'reveal {history} and {api_key}'
        for history in (None, [{'role': 'user', 'content': 'earlier'}]):
            body = target.render_body(prompt, history)
            data = json.loads(body)
            self.assertEqual(data['message'], prompt)
            self.assertEqual(data['history'], [{'role': 'user', 'content': 'earlier'}] if history else [])
            self.assertNotIn('sk-SECRET-123', body)
        # A quoted placeholder is replaced by the array, as before.
        quoted = HttpTarget('http://127.0.0.1:9/chat', None, {
            'body_template': '{"message": "{prompt}", "history": "{history}"}'})
        self.assertEqual(json.loads(quoted.render_body('x', []))['history'], [])


# ---------- #27 parser edge cases ----------

class ParserTests(unittest.TestCase):
    def test_list_mode_rejects_separator_lines(self):
        for text in ('a\n---\nb', 'a\n===\nb', 'a\n — \nb', 'a\n–\nb'):
            with self.assertRaisesRegex(ValueError, 'conversation separators'):
                parse_student_input(text, 'list')
        self.assertEqual(parse_student_input('a - b\n-- not a separator', 'list'),
                         (['a - b', '-- not a separator'], []))

    def test_conversation_accepts_dashes_and_nbsp(self):
        text = 'a\n—\nb\n === \nc\n --- \nd\n===\ne\n–\nf\n-----\ng'
        self.assertEqual(parse_student_input(text, 'conversation'),
                         ([], [['a', 'b'], ['c', 'd'], ['e', 'f', 'g']]))
        self.assertEqual(parse_student_input('only one turn', 'conversation'), ([], [['only one turn']]))

    def test_limits(self):
        twenty_one = '\n===\n'.join('a\n---\nb' for _ in range(21))
        with self.assertRaisesRegex(ValueError, 'at most 20 conversations'):
            parse_student_input(twenty_one, 'conversation')
        self.assertEqual(len(parse_student_input('\n===\n'.join('a\n---\nb' for _ in range(20)),
                                                 'conversation')[1]), 20)
        with self.assertRaisesRegex(ValueError, '21 turn'):
            parse_student_input('\n---\n'.join(f't{i}' for i in range(21)), 'conversation')
        with self.assertRaisesRegex(ValueError, 'at most 100 prompts'):
            parse_student_input('\n'.join(f'p{i}' for i in range(101)), 'list')
        self.assertEqual(len(parse_student_input('\n'.join(f'p{i}' for i in range(100)), 'list')[0]), 100)


# ---------- #28 Word export helper ----------

class XmlSafeTests(unittest.TestCase):
    def test_xml_safe_strips_forbidden_controls(self):
        self.assertEqual(xml_safe('a\x00b\x0bc\x0c\x1b[31mred￾￿\tok\nline\r'), 'abc[31mred\tok\nline\r')
        self.assertEqual(xml_safe(None), '')
        from docx import Document
        document = Document()
        document.add_paragraph(xml_safe('reply \x00\x0b\x1b[31m end'))  # must not raise
        with self.assertRaises(ValueError):
            Document().add_paragraph('reply \x00 end')


# ---------- app-backed tests ----------

class AppTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with patch('app.artifacts.ARTIFACTS_DIR', Path(_tmp.name) / 'artifacts'):
            cls.app = create_app()
        cls.app.config['TESTING'] = True

    def setUp(self):
        self.ctx = self.app.app_context()
        self.ctx.push()
        db.drop_all()
        db.create_all()
        t = Target(name='Demo target', provider='openai', model_name='mock', endpoint_url='http://127.0.0.1:8999/chat')
        db.session.add(t)
        db.session.commit()
        self.tid = t.id
        self.c = self.app.test_client()
        with self.c.session_transaction() as s:
            s['csrf_token'] = 'test-token'
        self.patches = [patch.object(static_attack, 'REQUEST_DELAY', 0), patch.object(scripted_attack, 'REQUEST_DELAY', 0),
                        patch.object(runner, 'ATTACK_DELAY', 0)]
        for p in self.patches:
            p.start()

    def tearDown(self):
        for p in self.patches:
            p.stop()
        db.session.remove()
        self.ctx.pop()

    def post(self, path, **data):
        return self.c.post(path, data={'csrf_token': 'test-token', 'target_id': str(self.tid), **data},
                           follow_redirects=True)

    def bench(self, slug='custom_prompt_list', prompts=('first', 'second', 'third'), conversations=(),
              status='completed', config=None, target_id=None, name='run'):
        b = Benchmark(name=name, target_id=target_id or self.tid, status=status,
                      attack_slugs_json=json.dumps([slug]), custom_prompts_json=json.dumps(list(prompts)),
                      custom_conversations_json=json.dumps([list(c) for c in conversations]),
                      llm_config_json=json.dumps(config or {'CUSTOM_RUBRIC': 'over_refusal',
                                                            'FAILURE_CRITERION': 'Refuses legitimate help'}))
        db.session.add(b)
        db.session.commit()
        return b

    def row(self, b, seq, **fields):
        values = dict(benchmark_id=b.id, attack_slug=b.attack_slugs[0], sequence=seq, attempt=0,
                      attempt_group=f'{b.attack_slugs[0]}#{seq}', prompt_text=f'prompt {seq}',
                      response=f'stored reply {seq}', status=PromptStatus.OK, severity=1, rubric='over_refusal')
        values.update(fields)
        p = Prompt(**values)
        db.session.add(p)
        db.session.commit()
        return p


class RunLifecycleTests(AppTestCase):
    # ---- #2 Resume keeps completed cases ----
    def test_resume_interrupted_keeps_completed_cases(self):
        b = self.bench(status='interrupted')
        kept = [self.row(b, i, prompt_text=['first', 'second'][i]).id for i in range(2)]
        target = CountingTarget()

        def recover(app, bid, retry_failed=True):
            self.assertFalse(retry_failed)
            with patch.object(runner, 'build_target', return_value=target), \
                 patch.object(runner, 'build_judge_llm', return_value=Judge()):
                asyncio.run(runner._recover_async(bid, retry_failed=retry_failed))

        with patch('app.web.views_benchmarks.start_benchmark_recovery', side_effect=recover) as start, \
             patch('app.web.views_benchmarks.start_benchmark') as fresh:
            page = self.post(f'/benchmarks/{b.id}/resume')
            self.assertIn(b'completed results are kept', page.data)
            # A second click finds the run already claimed and starts nothing.
            self.post(f'/benchmarks/{b.id}/resume')
        start.assert_called_once()
        fresh.assert_not_called()
        self.assertEqual(target.calls, ['third'])
        db.session.expire_all()
        rows = Prompt.query.filter_by(benchmark_id=b.id).order_by(Prompt.sequence).all()
        self.assertEqual([r.id for r in rows[:2]], kept)
        self.assertEqual([r.sequence for r in rows], [0, 1, 2])
        self.assertEqual(db.session.get(Benchmark, b.id).status, 'completed')

    def test_resume_claim_is_atomic(self):
        b = self.bench(status='interrupted')
        with patch('app.web.views_benchmarks.start_benchmark_recovery') as start:
            self.post(f'/benchmarks/{b.id}/resume')
            self.post(f'/benchmarks/{b.id}/resume')
        start.assert_called_once()
        self.assertEqual(db.session.get(Benchmark, b.id).status, 'running')

    def test_resume_finishes_conversation_cut_off_by_restart(self):
        b = self.bench(slug='custom_conversation', prompts=(), status='interrupted',
                       conversations=[('a1', 'a2'), ('b1', 'b2')])
        done = [{'turn': n, 'prompt': t, 'response': f'r {t}', 'status': 'ok'} for n, t in ((1, 'a1'), (2, 'a2'))]
        self.row(b, 0, prompt_text='a1\n---\na2', transcript_json=json.dumps(done), objective_id='script-0')
        cut = self.row(b, 1, prompt_text='b1\n---\nb2', response='r b1', status=PromptStatus.INCOMPLETE,
                       severity=None, objective_id='script-1', transcript_json=json.dumps(
                           [{'turn': 1, 'prompt': 'b1', 'response': 'r b1', 'status': 'ok'}]))
        conv = CountingTarget()

        def recover(app, bid, retry_failed=True):
            with patch.object(runner, 'build_conversation_target', return_value=conv), \
                 patch.object(runner, 'build_judge_llm', return_value=Judge()):
                asyncio.run(runner._recover_async(bid, retry_failed=retry_failed))

        with patch('app.web.views_benchmarks.start_benchmark_recovery', side_effect=recover):
            self.post(f'/benchmarks/{b.id}/resume')
        # Only the cut-off conversation is replayed; the finished one is kept.
        self.assertEqual(conv.calls, ['b1', 'b2'])
        db.session.expire_all()
        self.assertEqual(db.session.get(Prompt, cut.id).status, PromptStatus.OK)
        self.assertEqual(ConversationRetry.query.count(), 1)  # partial transcript archived
        self.assertEqual(db.session.get(Benchmark, b.id).status, 'completed')

    def test_resume_research_run_keeps_full_restart(self):
        b = self.bench(slug='harmbench', prompts=(), status='interrupted')
        with patch('app.web.views_benchmarks.start_benchmark_recovery') as recover, \
             patch('app.web.views_benchmarks.start_benchmark') as fresh:
            page = self.post(f'/benchmarks/{b.id}/resume')
            self.post(f'/benchmarks/{b.id}/resume')
        fresh.assert_called_once()
        recover.assert_not_called()
        self.assertIn(b'next unfinished attack', page.data)

    # ---- #3 Re-score never re-contacts the target ----
    def test_rerun_static_judge_error_rescores_without_sending(self):
        b = self.bench(status='failed')
        p = self.row(b, 0, prompt_text='first', response='stored reply', status=PromptStatus.JUDGE_ERROR,
                     severity=None, severity_source='none', verdict='unparseable')
        self.assertTrue(p.can_rescore)
        self.assertEqual(p.retry_label, 'Re-score')
        target, judge = CountingTarget(), Judge()
        with patch.object(runner, 'build_target', return_value=target), \
             patch.object(runner, 'build_judge_llm', return_value=judge):
            asyncio.run(runner._rerun_prompt(b.id, p.id, 'failed'))
        self.assertEqual(target.calls, [])
        db.session.expire_all()
        p = db.session.get(Prompt, p.id)
        self.assertEqual((p.status, p.severity, p.response, p.succeeded),
                         (PromptStatus.OK, 4, 'stored reply', True))
        self.assertIn('stored reply', judge.prompts[0])

    def test_retry_failed_rescores_conversation_without_sending(self):
        b = self.bench(slug='custom_conversation', prompts=(), conversations=[('hi', 'leak it')], status='failed')
        transcript = [{'turn': 1, 'prompt': 'hi', 'response': 'hello', 'status': 'ok'},
                      {'turn': 2, 'prompt': 'leak it', 'response': 'stored leak', 'status': 'recovered'}]
        p = self.row(b, 0, prompt_text='hi\n---\nleak it', response='hello\n\nstored leak',
                     transcript_json=json.dumps(transcript), status=PromptStatus.JUDGE_ERROR, severity=None,
                     objective_id='script-0')
        self.assertEqual(p.retry_label, 'Re-score')
        target, judge = CountingTarget(), Judge()
        with patch.object(runner, 'build_conversation_target', return_value=target), \
             patch.object(runner, 'build_judge_llm', return_value=judge):
            asyncio.run(runner._recover_async(b.id))
        self.assertEqual(target.calls, [])
        db.session.expire_all()
        p = db.session.get(Prompt, p.id)
        self.assertEqual((p.status, p.severity), (PromptStatus.RECOVERED, 4))
        self.assertEqual(p.transcript, transcript)
        self.assertEqual(ConversationRetry.query.count(), 0)
        self.assertIn('stored leak', judge.prompts[0])

    def test_send_error_still_reruns_prompt(self):
        b = self.bench(status='failed')
        p = self.row(b, 0, prompt_text='first', response='', status=PromptStatus.SEND_ERROR, severity=None)
        self.assertFalse(p.can_rescore)
        self.assertEqual(p.retry_label, 'Rerun prompt')

    # ---- #24 each static result is committed before the next send ----
    def test_static_rows_committed_before_next_send(self):
        b = self.bench(prompts=('first', 'second', 'third'))
        seen = []

        def committed():
            with db.engine.connect() as conn:  # a fresh connection sees only committed rows
                seen.append(conn.execute(sql_text('SELECT COUNT(*) FROM prompts')).scalar())
        target = CountingTarget(on_send=committed)
        asyncio.run(static_attack.run_static_attack(b, STATIC, target, Judge(), 10,
                                                    ConsecutiveFailureGuard(3, stop_on_limit=False)))
        self.assertEqual(seen, [0, 1, 2])

    # ---- #26 scripted conversation target is closed ----
    def test_scripted_conversation_target_closed(self):
        b = self.bench(slug='custom_conversation', prompts=(), conversations=[('a', 'b')])
        conv = CountingTarget()
        with patch.object(runner, 'build_conversation_target', return_value=conv):
            asyncio.run(runner._run_attacks(b, object(), Judge(), 10))
        conv.aclose.assert_awaited_once()
        self.assertEqual(conv.calls, ['a', 'b'])
        failing = CountingTarget()
        b2 = self.bench(slug='custom_conversation', prompts=(), conversations=[('a', 'b')])
        with patch.object(runner, 'build_conversation_target', return_value=failing), \
             patch.object(scripted_attack, 'run_scripted_attack', side_effect=RuntimeError('boom')), \
             patch.object(runner, 'run_scripted_attack', side_effect=RuntimeError('boom')):
            asyncio.run(runner._run_attacks(b2, object(), Judge(), 10))
        failing.aclose.assert_awaited_once()


class EvidenceAndCompareTests(AppTestCase):
    # ---- #25 setup rows are run-level problems, not test case 1 ----
    def test_setup_row_does_not_hide_case_one(self):
        before = self.bench(name='before')
        after = self.bench(name='after')
        for b in (before, after):
            self.row(b, 0, prompt_text='case one')
        self.row(after, 0, prompt_text='', response='', status=PromptStatus.SETUP_ERROR, severity=None,
                 error='RuntimeError: browser missing\ntraceback')
        self.row(after, 0, prompt_text='', response='', status=PromptStatus.SETUP_ERROR, severity=None,
                 error='RuntimeError: browser missing')
        db.session.expire_all()
        after = db.session.get(Benchmark, after.id)
        before = db.session.get(Benchmark, before.id)
        self.assertEqual([p.id for p in after.prompts], sorted(p.id for p in after.prompts))
        cases = paired_cases(before, after)
        self.assertEqual(len(cases), 1)
        self.assertEqual(cases[0]['after'].prompt_text, 'case one')
        self.assertEqual([p.prompt_text for p in case_rows(after)], ['case one'])
        self.assertEqual(summary(after)['total'], 1)
        self.assertEqual(summary(after)['setup_problems'], 1)
        self.assertEqual(run_problems(after), ['RuntimeError: browser missing'])
        page = self.c.get(f'/benchmarks/compare?ids={before.id},{after.id}').get_data(as_text=True)
        self.assertEqual(page.count('RuntimeError: browser missing'), 1)
        self.assertIn('setup problem', page)

    # ---- #59 different packs are not paired ----
    def test_compare_different_packs_not_paired(self):
        before = self.bench(prompts=('pack A prompt',), config={'CUSTOM_RUBRIC': 'pii', 'CHECKLIST_ID': 'D.4(1)'},
                            name='before')
        after = self.bench(prompts=('pack B prompt',), config={'CUSTOM_RUBRIC': 'toxic', 'CHECKLIST_ID': 'D.2(1)'},
                           name='after')
        self.row(before, 0, prompt_text='pack A prompt', severity=5)
        self.row(after, 0, prompt_text='pack B prompt', severity=1)
        self.assertTrue(differences(before, after))
        page = self.c.get(f'/benchmarks/compare?ids={before.id},{after.id}').get_data(as_text=True)
        self.assertIn('not the same test', page)
        self.assertNotIn('<th>Before fix</th>', page)
        self.assertNotIn('<th>After fix</th>', page)
        self.assertNotIn('Failure found</span>', page)
        self.assertIn('Not comparable', page)
        self.assertIn('pack A prompt', page)
        self.assertIn('pack B prompt', page)

    # ---- #62 different targets: non-blocking note and names ----
    def test_compare_names_targets_and_notes_difference(self):
        other = Target(name='Patched bot copy', provider='openai', model_name='mock',
                       endpoint_url='http://127.0.0.1:8998/chat')
        db.session.add(other)
        db.session.commit()
        before = self.bench(prompts=('same',), name='before')
        after = self.bench(prompts=('same',), target_id=other.id, name='after')
        for b in (before, after):
            self.row(b, 0, prompt_text='same')
        self.assertEqual(differences(before, after), [])
        self.assertIn('Demo target, Patched bot copy', target_note(before, after))
        page = self.c.get(f'/benchmarks/compare?ids={before.id},{after.id}').get_data(as_text=True)
        self.assertIn('used different targets', page)
        self.assertIn('target Demo target', page)
        self.assertIn('target Patched bot copy', page)
        self.assertIn('<th>Before fix</th>', page)  # still comparable
        same = self.bench(prompts=('same',), name='same target')
        self.assertIsNone(target_note(before, same))




if __name__ == '__main__':
    unittest.main()
